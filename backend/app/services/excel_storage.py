import os
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional
import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

# Path to users excel file
DATA_DIR = Path(__file__).resolve().parent.parent.parent / "data"
EXCEL_PATH = DATA_DIR / "users.xlsx"

HEADERS = [
    "ID",
    "Role",
    "Full Name",
    "Email",
    "Password",
    "Primary Sport",
    "Gender",
    "Date of Birth",
    "Location",
    "Experience Level",
    "Playing Position",
    "Created At",
]


def ensure_excel_file() -> None:
    """Ensure data directory and users.xlsx exist with appropriate header row."""
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    if not EXCEL_PATH.exists():
        wb = openpyxl.Workbook()
        ws = wb.active
        assert isinstance(ws, Worksheet)
        ws.title = "Users"
        ws.append(HEADERS)
        
        # Style headers
        for cell in ws[1]:
            cell.font = Font(bold=True)
            
        wb.save(EXCEL_PATH)


def get_all_users() -> List[Dict[str, str]]:
    """Retrieve all users stored in the Excel file as a list of dicts."""
    ensure_excel_file()
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb.active
    assert isinstance(ws, Worksheet)

    users: List[Dict[str, str]] = []
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) <= 1:
        return users

    headers = [str(h) if h else "" for h in rows[0]]

    for r in rows[1:]:
        if not any(r):
            continue
        user_dict = {}
        for idx, header in enumerate(headers):
            val = r[idx] if idx < len(r) and r[idx] is not None else ""
            user_dict[header] = str(val)
        users.append(user_dict)

    return users


def get_user_by_email(email: str) -> Optional[Dict[str, str]]:
    """Find a user in users.xlsx by email address (case-insensitive)."""
    email_clean = email.strip().lower()
    for user in get_all_users():
        if user.get("Email", "").strip().lower() == email_clean:
            return user
    return None


def save_user_to_excel(user_data: Dict[str, str]) -> Dict[str, str]:
    """Append a new user row to users.xlsx.
    
    Raises ValueError if email already exists.
    """
    ensure_excel_file()
    email_clean = user_data.get("email", "").strip().lower()
    
    if get_user_by_email(email_clean):
        raise ValueError("An account with this email already exists.")

    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active
    assert isinstance(ws, Worksheet)
    
    next_id = ws.max_row  # header is row 1, so max_row gives total count + 1 for next id
    created_at = datetime.utcnow().isoformat() + "Z"
    
    row_values = [
        next_id,
        user_data.get("role", "PLAYER").upper(),
        user_data.get("full_name", ""),
        email_clean,
        user_data.get("password", ""),
user_data.get("primary_sport", "Running / Sprinting"),
        user_data.get("gender", ""),
        user_data.get("dob", ""),
        user_data.get("location", ""),
        user_data.get("experience", "Beginner"),
        user_data.get("position", ""),
        created_at,
    ]
    
    ws.append(row_values)
    wb.save(EXCEL_PATH)
    
    return {
        "id": str(next_id),
        "role": user_data.get("role", "PLAYER").upper(),
        "full_name": user_data.get("full_name", ""),
        "email": email_clean,
        "primary_sport": user_data.get("primary_sport", "Running / Sprinting"),
        "gender": user_data.get("gender", ""),
        "dob": user_data.get("dob", ""),
        "location": user_data.get("location", ""),
        "experience": user_data.get("experience", "Beginner"),
        "position": user_data.get("position", ""),
        "created_at": created_at,
    }
